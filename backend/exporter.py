import copy, io, math, re, subprocess, unicodedata, zipfile
from functools import lru_cache
from .export_images import image_data,add_image
from .part_labels import part_display_label
import pypandoc
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
NS={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
@lru_cache(maxsize=512)
def _math_bytes(latex):
    if not latex.strip() or len(latex)>10000: raise ValueError('公式为空或过长')
    proc=subprocess.run([pypandoc.get_pandoc_path(),'-f','markdown+tex_math_dollars','-t','docx'],input=('$'+latex+'$').encode(),capture_output=True,timeout=20)
    if proc.returncode or proc.stderr: raise ValueError('公式转换警告: '+proc.stderr.decode()[:300])
    with zipfile.ZipFile(io.BytesIO(proc.stdout)) as z: root=etree.fromstring(z.read('word/document.xml'))
    nodes=root.xpath('//m:oMath',namespaces=NS)
    if len(nodes)!=1: raise ValueError('公式未生成 OMML: '+latex)
    return etree.tostring(nodes[0])
def math_xml(latex):
    from .text_format import math_notation
    latex=math_notation(latex)
    return etree.fromstring(_math_bytes(latex))
def spans(p,items,source=None):
    for s in items:
        if s['kind']=='image':
            data=image_data(source,s['text']);p.add_run().add_picture(io.BytesIO(data[0]),width=Pt(s.get('width') or 12),height=Pt(s.get('height') or 12))
        elif s['kind']=='math': p._p.append(math_xml(s['text']))
        else: p.add_run(s['text'])
def geometry(p,shapes):
    # Native VML objects; Mac Word editing remains a separate manual gate.
    v='urn:schemas-microsoft-com:vml'; w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    pict=etree.SubElement(p.add_run()._r,'{'+w+'}pict')
    group=etree.SubElement(pict,'{'+v+'}group',nsmap={'v':v}); group.set('style','width:300pt;height:180pt'); group.set('coordsize','300,180')
    for i,s in enumerate(shapes):
        kind=s['kind']; x=s['x']; y=s['y']; r=s['radius']
        tag='line' if kind=='line' else ('oval' if kind in ['point','circle'] else 'rect')
        n=etree.SubElement(group,'{'+v+'}'+tag); n.set('id','shape'+str(i)); n.set('strokecolor','#222222')
        if kind=='line': n.set('from',f'{x},{y}'); n.set('to',f'{s["x2"]},{s["y2"]}')
        else:
            radius=2 if kind=='point' else r
            n.set('style',f'position:absolute;left:{x-radius if kind!="label" else x};top:{y-radius if kind!="label" else y};width:{2*radius if kind!="label" else 45};height:{2*radius if kind!="label" else 20}')
            n.set('filled','t' if kind=='point' else 'f')
        if kind=='label':
            n.set('stroked','f'); box=etree.SubElement(n,'{'+v+'}textbox'); content=etree.SubElement(box,'{'+w+'}txbxContent'); pp=etree.SubElement(content,'{'+w+'}p'); rr=etree.SubElement(pp,'{'+w+'}r'); etree.SubElement(rr,'{'+w+'}t').text=s['label']
def blocks(doc,items,draft,source=None):
    image_paragraph=None;used_width=0
    sec=doc.sections[-1];available=(sec.page_width-sec.left_margin-sec.right_margin)/360000
    for b in items:
        if b['kind']!='image':image_paragraph=None;used_width=0
        if b['kind']=='paragraph': spans(doc.add_paragraph(),b['spans'],source)
        elif b['kind']=='equation': doc.add_paragraph()._p.append(math_xml(b['latex']))
        elif b['kind']=='table':
            from .table_layout import word_table
            word_table(doc,b['rows'],available,lambda p,items:spans(p,items,source))
        elif b['kind']=='geometry':
            if not draft: raise ValueError('Word 几何对象实际编辑尚待 Mac Word 人工验收，请使用草稿')
            p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); geometry(p,b['shapes'])
        elif b['kind']=='image':
            data=image_data(source,b['asset']);w,h=data[1]
            width=min(available,8 if w/h>2 else 5);height=width*h/w
            if height>4.5:width*=4.5/height;height=4.5
            if image_paragraph is None or used_width+width+.4>available:
                image_paragraph=doc.add_paragraph();used_width=0
            elif used_width:image_paragraph.add_run('  ')
            add_image(image_paragraph,data,width,height);used_width+=width+.4
        else: raise ValueError('不支持的题目内容类型')
def answer_blocks(q):
    return q.get('human_answer',[]) if q.get('answer_is_complete') or q.get('answer_edited') else (q.get('human_answer') or q.get('ai_answer') or q.get('original_answer',[]))
def explanation_blocks(q):
    return (q.get('human_explanation',[]) if q.get('explanation_edited') else (q.get('human_explanation') or q.get('ai_explanation') or q.get('original_explanation',[])))
def _question_height(nodes,doc):
    """Conservative point estimate used only to retain short-question grouping."""
    from docx.text.paragraph import Paragraph
    normal=doc.styles['Normal'];size=normal.font.size.pt
    width=(doc.sections[0].page_width-doc.sections[0].left_margin-doc.sections[0].right_margin)/12700
    def paragraph_height(element,available):
        paragraph=Paragraph(element,doc);fmt=paragraph.paragraph_format
        text=''.join('\n' if part.tag==qn('w:br') else (part.text or '') for part in element.xpath('.//w:t | .//m:t | .//w:br'))
        lines=sum(max(1,math.ceil(sum(1 if unicodedata.east_asian_width(char) in ('W','F') else .6 for char in line)*size/max(size,available))) for line in text.split('\n'))
        spacing=fmt.line_spacing or normal.paragraph_format.line_spacing or 1
        line_height=spacing.pt if hasattr(spacing,'pt') else size*float(spacing)*1.15
        drawings=element.xpath('.//wp:extent')
        height=max(lines*line_height,max((int(n.get('cy','0'))/12700+size for n in drawings),default=0))
        if element.xpath('.//w:pict'):height=max(height,180)
        before=fmt.space_before if fmt.space_before is not None else normal.paragraph_format.space_before
        after=fmt.space_after if fmt.space_after is not None else normal.paragraph_format.space_after
        return height+(before.pt if before else 0)+(after.pt if after else 0)
    height=0
    for node in nodes:
        if node.tag==qn('w:p'):height+=paragraph_height(node,width)
        elif node.tag==qn('w:tbl'):
            for row in node.findall(qn('w:tr')):
                cells=row.findall(qn('w:tc'));heights=[]
                for cell in cells:
                    cell_width=cell.find('./'+qn('w:tcPr')+'/'+qn('w:tcW'))
                    available=float(cell_width.get(qn('w:w')))/20 if cell_width is not None else width/max(1,len(cells))
                    heights.append(sum(paragraph_height(p,available-8) for p in cell.findall(qn('w:p')))+8)
                height+=max(heights,default=0)
    return height
def export_doc(questions,params,path):
    from .text_format import format_question_content
    questions=[dict(item,question=format_question_content(item['question'])) for item in questions]
    doc=Document(); sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    for x in ['top_margin','bottom_margin','left_margin','right_margin']: setattr(sec,x,Cm(float(params.get('margin',2))))
    for style in doc.styles:
        if style.type==1:
            style.font.name=params.get('latin_font',params.get('font','Times New Roman')); style.font.color.rgb=RGBColor(0,0,0)
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),params.get('chinese_font',params.get('font','SimSun')))
            for border in style._element.xpath('.//w:pBdr'): border.getparent().remove(border)
    normal=doc.styles['Normal']; normal.font.name=params.get('latin_font',params.get('font','Times New Roman')); normal.font.size=Pt(float(params.get('size',9))); normal.paragraph_format.line_spacing=float(params.get('line_spacing',1.25)); normal.paragraph_format.space_before=Pt(0); normal.paragraph_format.space_after=Pt(4.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),params.get('chinese_font',params.get('font','SimSun')))
    math_pr=doc.settings.element.find(qn('m:mathPr'))
    if math_pr is None:
        math_pr=OxmlElement('m:mathPr');doc.settings.element.append(math_pr)
    math_font=math_pr.find(qn('m:mathFont'))
    if math_font is None:
        math_font=OxmlElement('m:mathFont');math_pr.append(math_font)
    math_font.set(qn('m:val'),params.get('math_font','Cambria Math'))
    draft=params.get('draft',False); doc.add_heading(params.get('title','高中数学验证卷')+('（草稿）' if draft else ''),0)
    if params.get('kind')=='lesson':
        for k,label in [('objectives','教学目标'),('focus','重难点'),('process','教学过程'),('homework','作业')]: doc.add_heading(label,1); doc.add_paragraph(params.get(k,''))
    answers=[]
    for i,item in enumerate(questions):
        q=item['question']; first_node=len(doc._element.body); p=doc.add_paragraph(f'{i+1}. '+(f'（{item.get("score",5)}分）' if params.get('show_score',True) else '')); p.paragraph_format.keep_with_next=True
        if i:p.paragraph_format.space_before=Pt(8)
        stem=q['stem']
        # The export supplies its own number and score. Ignore only an isolated
        # source heading, never a heading that also contains question content.
        if stem and stem[0]['kind']=='paragraph' and all(span['kind']=='text' for span in stem[0]['spans']):
            source_heading=''.join(span['text'] for span in stem[0]['spans'])
            if re.fullmatch(r'\s*\d+\s*[.．]\s*[（(]\s*本(?:小)?题\s*满分\s*\d+(?:\.\d+)?\s*分\s*[）)]\s*',source_heading):
                stem=stem[1:]
        if stem and stem[0]['kind']=='paragraph':
            spans(p,stem[0]['spans'],q.get('source'));stem=stem[1:]
        blocks(doc,stem,draft,q.get('source'))
        for part in q['options']+q['subquestions']:
            content=part['blocks'];display_label=part_display_label(part)
            label=doc.add_paragraph(display_label+'  ' if display_label else '')
            if content and content[0]['kind']=='paragraph':
                spans(label,content[0]['spans'],q.get('source'));content=content[1:]
            else:label.paragraph_format.keep_with_next=True
            blocks(doc,content,draft,q.get('source'))
        from docx.text.paragraph import Paragraph
        question_nodes=list(doc._element.body)[first_node-1:-1]
        page_height=(sec.page_height-sec.top_margin-sec.bottom_margin)/12700
        # Leave room for the title on the first page and a small font-metric
        # allowance. Short questions retain their original whole-question keep.
        keep_whole=_question_height(question_nodes,doc)<page_height*.92-(72 if i==0 else 0)
        # Long questions must be allowed to flow across pages. Chaining every
        # paragraph (including all table cells) pushes a whole multi-page
        # question away from the title and creates near-empty pages.
        for node in question_nodes:
            if node.tag==qn('w:p'):
                paragraph=Paragraph(node,doc)
                paragraph.paragraph_format.keep_together=True
                paragraph.paragraph_format.widow_control=True
                text=paragraph.text.strip()
                if (text.startswith('【') and text.endswith('】')) or (len(text)<45 and text.endswith(('：',':'))):
                    paragraph.paragraph_format.keep_with_next=True
            elif node.tag==qn('w:tbl'):
                rows=node.findall(qn('w:tr'))
                for ri,row in enumerate(rows):
                    # Keep a source table together, but break the chain at its
                    # last row so the following question text remains free.
                    for element in row.xpath('.//w:p'):
                        paragraph=Paragraph(element,doc)
                        paragraph.paragraph_format.keep_with_next=ri<len(rows)-1
        if keep_whole:
            paragraphs=[Paragraph(p,doc) for node in question_nodes for p in ([node] if node.tag==qn('w:p') else node.xpath('.//w:p'))]
            for paragraph in paragraphs[:-1]:paragraph.paragraph_format.keep_with_next=True
            if paragraphs:paragraphs[-1].paragraph_format.keep_with_next=False
        for _ in range(min(30,max(0,int(item.get('blank',2))))): doc.add_paragraph(' ')
        if params.get('kind')!='student':
            if params.get('answer_position','end')=='inline':
                doc.add_paragraph('答案'); blocks(doc,answer_blocks(q),draft,q.get('source')); doc.add_paragraph('解析'); blocks(doc,explanation_blocks(q),draft,q.get('source'))
            else: answers.append((i,q))
    if answers:
        heading=doc.add_heading('答案与解析',1); heading.paragraph_format.page_break_before=True
        for i,q in answers:
            number=doc.add_paragraph(str(i+1)); number.paragraph_format.keep_with_next=True
            doc.add_paragraph('答案'); blocks(doc,answer_blocks(q),draft,q.get('source')); doc.add_paragraph('解析'); blocks(doc,explanation_blocks(q),draft,q.get('source'))
    p=sec.footer.paragraphs[0]; p.alignment=1; f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); p._p.append(f)
    # Remove template theme overrides and set explicit fonts on every text/math run.
    chinese=params.get('chinese_font',params.get('font','SimSun'))
    latin=params.get('latin_font',params.get('font','Times New Roman'))
    for fonts in doc.styles.element.xpath('.//w:rFonts'):
        for attr in list(fonts.attrib):
            if attr.endswith('Theme'): del fonts.attrib[attr]
        fonts.set(qn('w:ascii'),latin);fonts.set(qn('w:hAnsi'),latin);fonts.set(qn('w:eastAsia'),chinese)
    for run in doc.element.body.xpath('.//w:r | .//m:r'):
        rp=run.find(qn('w:rPr'))
        if rp is None: rp=OxmlElement('w:rPr');run.insert(0,rp)
        fonts=rp.find(qn('w:rFonts'))
        if fonts is None: fonts=OxmlElement('w:rFonts');rp.append(fonts)
        for attr in list(fonts.attrib):
            if attr.endswith('Theme'): del fonts.attrib[attr]
        family=params.get('math_font','Cambria Math') if run.tag==qn('m:r') else latin
        fonts.set(qn('w:ascii'),family);fonts.set(qn('w:hAnsi'),family);fonts.set(qn('w:eastAsia'),chinese)
        if run.tag==qn('m:r'):
            size=rp.find(qn('w:sz'))
            if size is None:size=OxmlElement('w:sz');rp.append(size)
            size.set(qn('w:val'),str(round(float(params.get('size',9))*2)))
    doc.save(path)

"""Editable, export-only paper snapshots. Never mutate source question records."""
import copy,json,uuid
from docx import Document
from docx.shared import Pt,Cm
from docx.oxml.ns import qn
from . import store
from .exporter import export_doc,math_xml,geometry,answer_blocks,explanation_blocks
from .part_labels import part_display_label
from .export_images import image_data,add_image


def paragraph(text='',**attrs):
    return dict(type='paragraph',attrs=attrs,content=[dict(type='text',text=text)] if text else [])


def inline(spans,source):
    out=[]
    for s in spans:
        if s['kind']=='math':out.append(dict(type='inlineMath',attrs=dict(latex=s['text'])))
        elif s['kind']=='image':out.append(dict(type='inlineSourceImage',attrs=dict(asset=s['text'],sourceId=source,width=s.get('width') or 12,height=s.get('height') or 12)))
        elif s['text']:out.append(dict(type='text',text=s['text']))
    return out


def nodes(blocks,source):
    out=[]
    for b in blocks:
        kind=b['kind']
        if kind=='paragraph':out.append(dict(type='paragraph',content=inline(b['spans'],source)))
        elif kind=='equation':out.append(dict(type='displayMath',attrs=dict(latex=b['latex'])))
        elif kind=='image':out.append(dict(type='sourceImage',attrs=dict(asset=b['asset'],sourceId=source)))
        elif kind=='geometry':out.append(dict(type='geometry',attrs=dict(shapes=b['shapes'])))
        elif kind=='table':out.append(dict(type='table',content=[dict(type='tableRow',content=[dict(type='tableCell',content=[dict(type='paragraph',content=inline(cell,source))]) for cell in row]) for row in b['rows']]))
    return out


def initial(snapshot):
    params=snapshot['params'];out=[paragraph(params.get('title','试卷'),align='center',fontSize=18)]
    if params.get('kind')=='lesson':
        for key,label in [('objectives','教学目标'),('focus','重难点'),('process','教学过程'),('homework','作业')]:out.extend([paragraph(label),paragraph(params.get(key,''))])
    answers=[]
    def solution(i,q):
        source=q.get('source',{}).get('id','')
        return [paragraph(f'{i+1}. 答案'),*nodes(answer_blocks(q),source),paragraph('解析'),*nodes(explanation_blocks(q),source)]
    for i,item in enumerate(snapshot['items']):
        q=item['question'];source=q.get('source',{}).get('id','');stem=nodes(q['stem'],source)
        prefix=f'{i+1}. '+(f'（{item.get("score",5)}分）' if params.get('show_score',True) else '')
        if stem and stem[0]['type']=='paragraph':stem[0]['content'].insert(0,dict(type='text',text=prefix))
        else:stem.insert(0,paragraph(prefix))
        out.extend(stem)
        for part in q['options']+q['subquestions']:
            bs=nodes(part['blocks'],source);label=part_display_label(part)
            if bs and bs[0]['type']=='paragraph':
                if label:bs[0]['content'].insert(0,dict(type='text',text=label+' '))
            else:bs.insert(0,paragraph(label))
            out.extend(bs)
        out.extend(paragraph() for _ in range(min(30,max(0,int(item.get('blank',2))))))
        if params.get('kind')!='student':
            if params.get('answer_position','end')=='inline':out.extend(solution(i,q))
            else:answers.extend(solution(i,q))
    if answers:out.extend([dict(type='paperBreak'),paragraph('答案与解析',align='center',fontSize=16),*answers])
    return dict(type='doc',content=out)


def get(export_id):
    snapshot=store.get('exports',export_id)
    if not snapshot:raise ValueError('导出文件不存在')
    return dict(id=export_id,document=snapshot.get('document') or initial(snapshot),params=snapshot['params'])


def write(document,params,snapshot,path):
    if not isinstance(document,dict) or document.get('type')!='doc':raise ValueError('文档格式错误')
    if len(json.dumps(document,ensure_ascii=False))>2000000:raise ValueError('文档内容过大')
    allowed_sources={item['question'].get('source',{}).get('id') for item in snapshot['items']}
    # Reuse the Word template, font settings, margins and footer from normal export.
    export_doc([],params,path);doc=Document(path)
    for child in list(doc.element.body):
        if child.tag!=qn('w:sectPr'):doc.element.body.remove(child)
    count=0
    def source(attrs):
        sid=attrs.get('sourceId')
        if sid not in allowed_sources:raise ValueError('图片不属于当前试卷的来源文件')
        return {'id':sid}
    def number(value,default,low,high):
        value=float(default if value is None else value)
        if not low<=value<=high:raise ValueError('文档排版参数超范围')
        return value
    def spans(p,children):
        for node in children:
            kind=node['type'];attrs=node.get('attrs') or {}
            if kind=='text':
                run=p.add_run(node['text'])
                for mark in node.get('marks',[]):
                    if mark['type'] not in ['bold','italic','underline','strike']:raise ValueError('不支持的文字格式')
                    setattr(run.font,mark['type'],True)
            elif kind=='hardBreak':p.add_run().add_break()
            elif kind=='inlineMath':p._p.append(math_xml(attrs['latex']))
            elif kind=='inlineSourceImage':
                data=image_data(source(attrs),attrs['asset']);p.add_run().add_picture(__import__('io').BytesIO(data[0]),width=Pt(number(attrs.get('width'),12,1,2000)),height=Pt(number(attrs.get('height'),12,1,2000)))
            else:raise ValueError('不支持的行内内容：'+kind)
    def render(parent,children):
        nonlocal count
        for node in children:
            count+=1
            if count>20000:raise ValueError('文档段落过多')
            kind=node['type'];attrs=node.get('attrs') or {};content=node.get('content') or []
            if kind=='paragraph':
                p=parent.add_paragraph();spans(p,content)
                align=attrs.get('align') or 'left'
                if align not in ['left','center','right','justify']:raise ValueError('段落对齐无效')
                p.alignment={'left':0,'center':1,'right':2,'justify':3}[align]
                size=number(attrs.get('fontSize'),params.get('size',9),8,48)
                p.paragraph_format.first_line_indent=Pt(size*number(attrs.get('indent'),0,0,4))
                p.paragraph_format.line_spacing=number(attrs.get('lineSpacing'),params.get('line_spacing',1.25),1,3)
                for run in p.runs:run.font.size=Pt(size)
                from docx.oxml import OxmlElement
                for run in p._p.xpath('.//m:r'):
                    rp=run.find(qn('w:rPr'))
                    if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
                    sz=rp.find(qn('w:sz'))
                    if sz is None:sz=OxmlElement('w:sz');rp.append(sz)
                    sz.set(qn('w:val'),str(round(size*2)))
            elif kind=='displayMath':parent.add_paragraph()._p.append(math_xml(attrs['latex']))
            elif kind=='paperBreak':parent.add_paragraph().paragraph_format.page_break_before=True
            elif kind=='sourceImage':
                data=image_data(source(attrs),attrs['asset']);available=21-2*float(params.get('margin',2));width=min(available,number(attrs.get('widthCm'),5,1,19));height=width*data[1][1]/data[1][0]
                add_image(parent.add_paragraph(),data,width,height)
            elif kind=='geometry':
                if not params.get('draft'):raise ValueError('几何对象需使用草稿导出')
                geometry(parent.add_paragraph(),attrs['shapes'])
            elif kind=='table':
                if not content or not content[0].get('content'):raise ValueError('表格不能为空')
                cols=len(content[0]['content']);table=parent.add_table(rows=len(content),cols=cols);table.style='Table Grid'
                for ri,row in enumerate(content):
                    if row['type']!='tableRow' or len(row['content'])!=cols:raise ValueError('表格必须为矩形')
                    for ci,cell in enumerate(row['content']):
                        if cell['type'] not in ['tableCell','tableHeader'] or any((cell.get('attrs') or {}).get(k,1)!=1 for k in ['colspan','rowspan']):raise ValueError('暂不支持合并单元格')
                        target=table.cell(ri,ci);render(target,cell.get('content',[]));first=target.paragraphs[0]
                        if len(target.paragraphs)>1 and not first.text:first._p.getparent().remove(first._p)
            else:raise ValueError('不支持的文档内容：'+kind)
    render(doc,document.get('content',[]))
    # Set explicit fonts for edited runs and editable OMML equations.
    for run in doc.element.body.xpath('.//w:r | .//m:r'):
        from docx.oxml import OxmlElement
        rp=run.find(qn('w:rPr'))
        if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
        fonts=rp.find(qn('w:rFonts'))
        if fonts is None:fonts=OxmlElement('w:rFonts');rp.append(fonts)
        family=params.get('math_font','Cambria Math') if run.tag==qn('m:r') else params.get('latin_font','Times New Roman')
        for k in ['ascii','hAnsi']:fonts.set(qn('w:'+k),family)
        fonts.set(qn('w:eastAsia'),params.get('chinese_font','SimSun'))
    doc.save(path)


def save(export_id,data):
    previous=store.get('exports',export_id)
    if not previous:raise ValueError('原导出文件不存在')
    params=copy.deepcopy(previous['params']);params.update({k:v for k,v in data.get('params',{}).items() if k in ['margin','size','line_spacing','chinese_font','latin_font','math_font']})
    if not 1<=float(params.get('margin',2))<=5 or not 8<=float(params.get('size',9))<=24:raise ValueError('页边距或字号超范围')
    identity=uuid.uuid4().hex;path=store.DATA/'exports'/(identity+'.docx')
    try:write(data['document'],params,previous,path)
    except Exception:
        path.unlink(missing_ok=True);raise
    snapshot=dict(previous,id=identity,document=data['document'],params=params,parent_export_id=export_id)
    store.put('exports',identity,snapshot)
    return dict(id=identity,url='/api/export/'+identity)

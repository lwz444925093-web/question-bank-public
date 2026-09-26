"""Content-weighted widths shared in principle by browser and Word rendering."""
import re

def cell_length(cell):
 text=''.join(s.get('text','') for s in cell)
 text=re.sub(r'\\[A-Za-z]+', 'x', text)
 return max((sum(1 if ord(c)>255 else .55 for c in line) for line in text.splitlines()),default=0)

def column_weights(rows):
 cols=max((len(r) for r in rows),default=0)
 return [max(4,min(18,max((cell_length(r[j]) for r in rows if j<len(r)),default=0)+2)) for j in range(cols)]

def word_table(doc,rows,available,write_spans):
 from docx.shared import Cm,Pt
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
 from docx.enum.text import WD_ALIGN_PARAGRAPH
 weights=column_weights(rows)
 if not weights:return
 table=doc.add_table(rows=len(rows),cols=len(weights));table.style='Table Grid';table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=[available*w/sum(weights) for w in weights]
 for col,width in zip(table.columns,widths):col.width=Cm(width)
 margins=OxmlElement('w:tblCellMar')
 for key,value in [('top',80),('bottom',80),('left',80),('right',80)]:
  node=OxmlElement('w:'+key);node.set(qn('w:w'),str(value));node.set(qn('w:type'),'dxa');margins.append(node)
 table._tbl.tblPr.append(margins)
 for i,row in enumerate(rows):
  trpr=table.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j in range(len(weights)):
   content=row[j] if j<len(row) else [];cell=table.cell(i,j);cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if cell_length(content)>12 else WD_ALIGN_PARAGRAPH.CENTER
   p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
   p.paragraph_format.keep_together=True
   write_spans(p,content)
 return table
